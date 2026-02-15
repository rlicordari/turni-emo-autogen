from __future__ import annotations

import calendar
import datetime as dt
from copy import copy
from pathlib import Path
from typing import Dict, Optional, Tuple

import openpyxl
from openpyxl.utils import column_index_from_string

from .config import ITALIAN_DOW, COL_J, COL_M, COL_N, COL_O, COL_P


def _load_style_ws(style_template_path: Path) -> openpyxl.worksheet.worksheet.Worksheet:
    wb = openpyxl.load_workbook(style_template_path)
    return wb[wb.sheetnames[0]]


def _weekday_style_rows(ws_style) -> Dict[int, int]:
    """Map weekday (0..6) -> example row index in style sheet."""
    mapping: Dict[int, int] = {}
    for r in range(2, ws_style.max_row + 1):
        v = ws_style.cell(row=r, column=1).value
        if isinstance(v, dt.datetime):
            d = v.date()
        elif isinstance(v, dt.date):
            d = v
        else:
            continue
        wd = d.weekday()
        if wd not in mapping:
            mapping[wd] = r
        if len(mapping) == 7:
            break
    return mapping


def _copy_cell_style(src, dst) -> None:
    dst._style = copy(src._style)
    dst.number_format = src.number_format
    dst.font = copy(src.font)
    dst.fill = copy(src.fill)
    dst.border = copy(src.border)
    dst.alignment = copy(src.alignment)
    dst.protection = copy(src.protection)


def create_month_workbook_from_style(style_template_path: Path, year: int, month: int) -> openpyxl.Workbook:
    """Create a fresh workbook for year/month, applying the look from Style_Template.xlsx."""
    ws_style = _load_style_ws(style_template_path)
    style_rows = _weekday_style_rows(ws_style)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"GUARDIE_{calendar.month_name[month].upper()}_{year}"

    # Copy merges (header / static parts)
    for merged in ws_style.merged_cells.ranges:
        ws.merge_cells(str(merged))

    # Copy column dimensions (widths)
    for col_letter, dim in ws_style.column_dimensions.items():
        ws.column_dimensions[col_letter].width = dim.width

    # Copy row dimensions for header row (and we'll copy day rows later)
    if 1 in ws_style.row_dimensions:
        ws.row_dimensions[1].height = ws_style.row_dimensions[1].height

    max_col = ws_style.max_column

    # Header row: values + style
    for c in range(1, max_col + 1):
        src = ws_style.cell(row=1, column=c)
        dst = ws.cell(row=1, column=c)
        dst.value = src.value
        _copy_cell_style(src, dst)

    # Fill dates + day label + styles for each day row
    last_day = calendar.monthrange(year, month)[1]
    for day in range(1, last_day + 1):
        d = dt.date(year, month, day)
        r = day + 1
        src_r = style_rows.get(d.weekday(), 2)

        # Copy entire row style (all columns)
        ws.row_dimensions[r].height = ws_style.row_dimensions[src_r].height
        for c in range(1, max_col + 1):
            src = ws_style.cell(row=src_r, column=c)
            dst = ws.cell(row=r, column=c)
            _copy_cell_style(src, dst)

        # Set date + weekday label (values)
        ws.cell(row=r, column=1).value = d
        ws.cell(row=r, column=1).number_format = "dd/mm/yyyy"
        ws.cell(row=r, column=2).value = ITALIAN_DOW[d.weekday()]

    # Nice-to-have view options
    try:
        ws.freeze_panes = "A2"
    except Exception:
        pass

    return wb


def write_assignments(ws, assignment_by_date: Dict[dt.date, Dict[str, str]]) -> None:
    """Write J/M/N/O/P into their columns, matching by date in column A."""
    # Build date->row lookup
    date_to_row: Dict[dt.date, int] = {}
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if isinstance(v, dt.datetime):
            d = v.date()
        elif isinstance(v, dt.date):
            d = v
        else:
            continue
        date_to_row[d] = r

    def col_idx(letter: str) -> int:
        return column_index_from_string(letter)

    for d, rowvals in assignment_by_date.items():
        r = date_to_row.get(d)
        if not r:
            continue
        # J
        ws.cell(row=r, column=col_idx(COL_J)).value = rowvals.get("J", "")
        # M N O P
        ws.cell(row=r, column=col_idx(COL_M)).value = rowvals.get("M", "")
        ws.cell(row=r, column=col_idx(COL_N)).value = rowvals.get("N", "")
        ws.cell(row=r, column=col_idx(COL_O)).value = rowvals.get("O", "")
        ws.cell(row=r, column=col_idx(COL_P)).value = rowvals.get("P", "")
